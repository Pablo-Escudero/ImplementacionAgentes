"""Extracción de texto de los archivos subidos (PDF, TXT, Markdown, CSV y DOCX)."""

import io
import os

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".markdown", ".csv", ".docx"}


class UnsupportedFormatError(Exception):
    pass


class ExtractionError(Exception):
    """El archivo no se pudo leer o no contiene texto utilizable."""


def get_extension(filename: str) -> str:
    return os.path.splitext(filename)[1].lower()


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("No se pudo decodificar el archivo de texto.")


def _extract_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("El PDF está protegido con contraseña.")
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("No se pudo leer el PDF (archivo corrupto o inválido).") from exc


def _extract_docx(data: bytes) -> str:
    try:
        doc = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("No se pudo leer el DOCX (archivo corrupto o inválido).") from exc
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    """Devuelve el texto del archivo o lanza UnsupportedFormatError / ExtractionError."""
    ext = get_extension(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFormatError(ext)
    if not data:
        raise ExtractionError("El archivo está vacío.")

    if ext == ".pdf":
        text = _extract_pdf(data)
    elif ext == ".docx":
        text = _extract_docx(data)
    else:
        text = _decode(data)

    # PostgreSQL no admite el carácter NUL en columnas de texto.
    text = text.replace("\x00", "").strip()
    if not text:
        raise ExtractionError(
            "No se encontró texto en el archivo (¿PDF escaneado como imagen o documento vacío?)."
        )
    return text
